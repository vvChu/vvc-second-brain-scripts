"""VvC Second Brain — Document Generation Worker (v7.0).

Generates Word (.docx) files from LLM markdown response.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from docx import Document

from core.config import cfg
from core.log import log
from services.diagram_base import spawn_worker

_logger = logging.getLogger("vvc.doc")


def trigger_doc_generation(filename: str, source_text: str) -> None:
    """Trigger background DOCX generation.

    Args:
        filename: Filename (e.g. "report.docx").
        source_text: Full response text to convert to word.
    """
    spawn_worker(
        target=_generate_doc,
        args=(filename, source_text),
        name=f"doc-{filename}",
    )


def _render_line_to_docx(doc: Document, line: str) -> None:
    """Render a single markdown line into the Document."""
    if line.startswith("> "):
        line = line[2:]
    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=3)
    elif line.startswith(("- ", "* ")):
        text = line[2:].strip().replace("**", "")
        doc.add_paragraph(text, style="List Bullet")
    elif re.match(r"^\d+\.\s+", line):
        text = re.sub(r"^\d+\.\s+", "", line).strip().replace("**", "")
        doc.add_paragraph(text, style="List Number")
    else:
        doc.add_paragraph(line.replace("**", ""))


def _render_markdown_to_docx(doc: Document, source_text: str) -> None:
    """Parse simplified markdown lines and append to docx Document."""
    clean_text = re.sub(r"!\[\[.*?\]\]", "", source_text)
    for raw_line in clean_text.split("\n"):
        line = raw_line.strip()
        if line:
            _render_line_to_docx(doc, line)


def _generate_doc(filename: str, source_text: str) -> None:
    """Worker function: generate DOCX file."""
    _logger.info(f"Generating DOCX: {filename}")
    log("diagram", f"DOCX generation started: {filename}")
    try:
        doc = Document()
        _render_markdown_to_docx(doc, source_text)
        output_dir = cfg.vault_root / "03 - Resources" / "doc"
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / filename
        doc.save(str(output_path))
        _logger.info(f"DOCX saved: {filename}")
        log("diagram", f"DOCX created: {filename}")
    except Exception as e:
        _logger.error(f"DOCX generation failed for {filename}: {e}")
        log("error", f"DOCX generation failed: {e}")
